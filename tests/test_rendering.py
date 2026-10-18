"""Pure formatting tests for to_markdown/to_docx -- fast, no model load."""

import io

from docx import Document

from app.transcriber import Segment, to_docx, to_markdown

SEGMENTS = [
    Segment(start=0.0, end=2.0, text="Hola, buenos dias."),
    Segment(start=2.0, end=4.5, text="Como estas hoy?"),
]


def test_markdown_with_timestamps_has_one_line_per_segment():
    markdown = to_markdown(
        "clip.wav", SEGMENTS, detected_language=None, with_timestamps=True
    )

    assert "**[00:00:00]** Hola, buenos dias." in markdown
    assert "**[00:00:02]** Como estas hoy?" in markdown


def test_markdown_without_timestamps_is_one_continuous_block():
    markdown = to_markdown(
        "clip.wav", SEGMENTS, detected_language=None, with_timestamps=False
    )

    assert "[00:00:" not in markdown
    assert "Hola, buenos dias. Como estas hoy?" in markdown


def test_docx_with_timestamps_has_one_paragraph_per_segment():
    content = to_docx(
        "clip.wav", SEGMENTS, detected_language=None, with_timestamps=True
    )
    doc = Document(io.BytesIO(content))

    body_paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert any(
        "Hola, buenos dias." in p and p.startswith("[00:00:00]")
        for p in body_paragraphs
    )
    assert any(
        "Como estas hoy?" in p and p.startswith("[00:00:02]") for p in body_paragraphs
    )


def test_docx_without_timestamps_is_one_continuous_paragraph():
    content = to_docx(
        "clip.wav", SEGMENTS, detected_language=None, with_timestamps=False
    )
    doc = Document(io.BytesIO(content))

    body_paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert any(p == "Hola, buenos dias. Como estas hoy?" for p in body_paragraphs)
